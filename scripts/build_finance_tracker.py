#!/usr/bin/env python3
"""Build Personal_Finance_Tracker_v2.xlsx.

Six tabs:
  1. Transactions  - the ONLY place you type numbers
  2. Dashboard     - visual overview, month picker
  3. Loan Payoff   - the ideal plan AND every payment's running balance, one tab
  4. Monthly       - month by month, category spending with a month picker, pay periods
  5. Savings       - where your savings go, you set the target split
  6. Start Here    - how to use it + the one-time setup numbers
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.chart import LineChart, BarChart, Reference

OUT = "/Users/shazilfarukh/Desktop/Quant_algo_trading/Personal_Finance_Tracker_v2.xlsx"

# ------------------------------------------------------------------ palette
INK = "16283C"      # near-black navy, body text
NAVY = "1F3A5F"     # banners
TEAL = "2F6F6B"
GREEN = "1E7A46"
RED = "B02A2A"
AMBER = "9A6410"
GREY = "6E7C8A"
LINE = "CBD5DF"

CARD = "F7FAFC"
BAND = "EDF2F7"
ZEBRA = "F9FAFB"
INPUT_FILL = "FFF6D8"
GOODF = "E3F4E9"
BADF = "FBE4E4"
WARNF = "FDF1DC"
WHITE = "FFFFFF"

MONEY = '$#,##0.00'
MONEY0 = '$#,##0'
PCT = '0.0%'
PCT2 = '0.00%'
DATEF = 'mm/dd/yyyy'
MONF = 'mmm yyyy'
INT = '0'
TEXTF = '@'

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
NOBOX = Border()

F_NOTE = Font(italic=True, size=9, color=GREY)
F_HDR = Font(bold=True, size=10, color=WHITE)
F_BANNER = Font(bold=True, size=11, color=WHITE)
F_LBL = Font(bold=True, size=9, color=GREY)
F_KPI = Font(bold=True, size=16, color=INK)
F_KPI_S = Font(bold=True, size=12, color=INK)
F_BODY = Font(size=10, color=INK)
F_BOLD = Font(bold=True, size=10, color=INK)
F_SECT = Font(bold=True, size=11, color=NAVY)

FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_BANNER = PatternFill("solid", fgColor=NAVY)
FILL_CARD = PatternFill("solid", fgColor=CARD)
FILL_BAND = PatternFill("solid", fgColor=BAND)
FILL_INPUT = PatternFill("solid", fgColor=INPUT_FILL)

N_TX = 1001
N_MONTHS = 36
N_PAY = 40                       # payment rows per loan ledger
START = dt.datetime(2026, 8, 1)
TARGET = dt.datetime(2028, 7, 31)

CATEGORIES = ["Income", "Housing", "Food", "Transportation", "Bills",
              "Student Debt", "Savings/Investments", "Shopping",
              "Health/Fitness", "Travel", "Family", "Education",
              "Entertainment", "Transfer", "Other"]
SPEND_CATS = [c for c in CATEGORIES if c != "Income"]
ACCOUNTS = ["Checking", "Savings", "Credit Card", "Cash", "Other"]
BUCKETS = ["Climb Launch", "Climb UAS", "Emergency Fund", "Brokerage / Stocks",
           "Masters Fund", "Roth IRA", "Crypto", "Other Savings"]
SAV_BUCKETS = BUCKETS[2:]

# ------------------------------------------------------------------ helpers


def sheet_title(ws, text, sub, last_col):
    lc = get_column_letter(last_col)
    ws.merge_cells(f"A1:{lc}1")
    c = ws["A1"]
    c.value = text
    c.font = Font(bold=True, size=17, color=WHITE)
    c.alignment = Alignment(vertical="center", indent=1)
    for i in range(1, last_col + 1):
        ws.cell(row=1, column=i).fill = FILL_HDR
    ws.row_dimensions[1].height = 32
    ws.merge_cells(f"A2:{lc}2")
    c = ws["A2"]
    c.value = sub
    c.font = F_NOTE
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 16


def banner(ws, row, c1, c2, text):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    c = ws.cell(row=row, column=c1, value=text)
    c.font = F_BANNER
    c.alignment = Alignment(vertical="center", indent=1)
    for col in range(c1, c2 + 1):
        ws.cell(row=row, column=col).fill = FILL_BANNER
    ws.row_dimensions[row].height = 21


def table_head(ws, row, col0, labels, widths=None, height=30):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=lab)
        c.font = F_HDR
        c.fill = PatternFill("solid", fgColor="3D5A80")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        if widths:
            ws.column_dimensions[get_column_letter(col0 + i)].width = widths[i]
    ws.row_dimensions[row].height = height


def cell(ws, row, col, value=None, fmt=None, font=F_BODY, fill=None,
         align=None, border=BOX, wrap=False):
    c = ws.cell(row=row, column=col)
    if value is not None:
        c.value = value
    if fmt:
        c.number_format = fmt
    c.font = font
    if fill:
        c.fill = fill
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    c.border = border
    return c


def kpi(ws, lrow, col, lab, formula, fmt=MONEY0, small=False, span=1):
    lc = ws.cell(row=lrow, column=col, value=lab.upper())
    lc.font = F_LBL
    lc.alignment = Alignment(horizontal="center")
    lc.fill = FILL_CARD
    if span > 1:
        ws.merge_cells(start_row=lrow, start_column=col, end_row=lrow, end_column=col + span - 1)
        ws.merge_cells(start_row=lrow + 1, start_column=col, end_row=lrow + 1, end_column=col + span - 1)
    v = ws.cell(row=lrow + 1, column=col, value=formula)
    v.number_format = fmt
    v.font = F_KPI_S if small else F_KPI
    v.fill = FILL_CARD
    v.alignment = Alignment(horizontal="center", vertical="center")
    for i in range(span):
        ws.cell(row=lrow, column=col + i).fill = FILL_CARD
        ws.cell(row=lrow + 1, column=col + i).fill = FILL_CARD
        ws.cell(row=lrow, column=col + i).border = Border(top=thin, left=thin, right=thin)
        ws.cell(row=lrow + 1, column=col + i).border = Border(bottom=thin, left=thin, right=thin)
    ws.row_dimensions[lrow].height = 15
    ws.row_dimensions[lrow + 1].height = 26
    return v


def zebra(ws, rng, first_row):
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'MOD(ROW()-{first_row},2)=1'],
        fill=PatternFill("solid", bgColor=ZEBRA), stopIfTrue=False))


def month_picker(ws, row, col, list_col, label_text="MONTH:"):
    """Label + validated month cell. Month list lives hidden in `list_col`."""
    lc = ws.cell(row=row, column=col - 1, value=label_text)
    lc.font = Font(bold=True, size=10, color=NAVY)
    lc.alignment = Alignment(horizontal="right", vertical="center")
    c = ws.cell(row=row, column=col, value=START)
    c.number_format = MONF
    c.font = Font(bold=True, size=12, color=INK)
    c.fill = FILL_INPUT
    c.border = Border(left=Side("medium", color=AMBER), right=Side("medium", color=AMBER),
                      top=Side("medium", color=AMBER), bottom=Side("medium", color=AMBER))
    c.alignment = Alignment(horizontal="center", vertical="center")
    L = get_column_letter(list_col)
    for i in range(N_MONTHS):
        r = 6 + i
        mc = ws.cell(row=r, column=list_col)
        mc.value = START if i == 0 else f"=EDATE({L}{r-1},1)"
        mc.number_format = MONF
    ws.column_dimensions[L].hidden = True
    dv = DataValidation(type="list", allow_blank=False,
                        formula1=f"${L}$6:${L}${5+N_MONTHS}")
    dv.prompt = "Pick a month"
    ws.add_data_validation(dv)
    dv.add(ws.cell(row=row, column=col).coordinate)
    return f"${get_column_letter(col)}${row}"


def bar_cell(ws, row, c1, c2, pct_ref, color=GREEN, n=34):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    c = ws.cell(row=row, column=c1)
    full, empty = "\u2588", "\u2591"
    c.value = (f'=REPT("{full}",ROUND(MIN(1,MAX(0,{pct_ref}))*{n},0))'
               f'&REPT("{empty}",{n}-ROUND(MIN(1,MAX(0,{pct_ref}))*{n},0))'
               f'&"  "&TEXT({pct_ref},"0.0%")')
    c.font = Font(bold=True, size=11, color=color)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for i in range(c1, c2 + 1):
        ws.cell(row=row, column=i).fill = FILL_CARD
    ws.row_dimensions[row].height = 20
    return c


# ------------------------------------------------------------------ workbook
wb = Workbook()
tx = wb.active
tx.title = "Transactions"
db = wb.create_sheet("Dashboard")
lp = wb.create_sheet("Loan Payoff")
mo = wb.create_sheet("Monthly")
sv = wb.create_sheet("Savings")
sh = wb.create_sheet("Start Here")

tx.sheet_properties.tabColor = "E0A80D"
db.sheet_properties.tabColor = "1F3A5F"
lp.sheet_properties.tabColor = "B02A2A"
mo.sheet_properties.tabColor = "2F6F6B"
sv.sheet_properties.tabColor = "1E7A46"
sh.sheet_properties.tabColor = "8794A1"

# cross-sheet references, named once so the formulas below stay readable
LP, SH, TXS = "'Loan Payoff'", "'Start Here'", "Transactions"
BAL_L, REQ_L = f"{LP}!$A$6", f"{LP}!$B$6"
BAL_U, REQ_U = f"{LP}!$C$6", f"{LP}!$D$6"
BAL_T, REQ_T = f"{LP}!$E$6", f"{LP}!$F$6"
TGT_D, PROJ_D = f"{LP}!$G$6", f"{LP}!$H$6"
PAID_L, INT_L, NPAY_L = f"{LP}!$K$49", f"{LP}!$K$50", f"{LP}!$K$51"
PAID_U, INT_U, NPAY_U = f"{LP}!$K$93", f"{LP}!$K$94", f"{LP}!$K$95"
OPEN_L, OPEN_U, OPEN_T = f"{SH}!$D$16", f"{SH}!$D$17", f"{SH}!$D$18"
APR_L, APR_U = f"{SH}!$E$16", f"{SH}!$E$17"
ASOF_L, ASOF_U = f"{SH}!$F$16", f"{SH}!$F$17"
TARGET_DATE, MONTHS_LEFT = f"{SH}!$B$21", f"{SH}!$B$22"
PLAN_INCOME, PLAN_SAVE, PLAN_LOAN = f"{SH}!$B$26", f"{SH}!$B$32", f"{SH}!$B$33"
PLAN_OUT, PLAN_LEFT = f"{SH}!$B$34", f"{SH}!$B$35"

# Data on Transactions runs from TX0 to TX1. Every shared range is built from
# these two so the two can never drift apart again.
TX0 = 5                          # first data row (row 4 is the header)
TX1 = TX0 + N_TX - 2             # 1004 -> 1000 usable rows
TXD = f"{TXS}!$A${TX0}:$A${TX1}"
TXC = f"{TXS}!$C${TX0}:$C${TX1}"
TXIN = f"{TXS}!$D${TX0}:$D${TX1}"
TXOUT = f"{TXS}!$E${TX0}:$E${TX1}"
TXG = f"{TXS}!$G${TX0}:$G${TX1}"


# Transfers between your own accounts are not income and not spending, so every
# cash-flow figure in the workbook filters them out.
NT = f',{TXC},"<>Transfer"'


def m_in(sel, extra=NT):
    """SUMIFS of Money In for the month starting at `sel`."""
    return (f'SUMIFS({TXIN},{TXD},">="&{sel},{TXD},"<"&EDATE({sel},1){extra})')


def m_out(sel, extra=NT):
    return (f'SUMIFS({TXOUT},{TXD},">="&{sel},{TXD},"<"&EDATE({sel},1){extra})')


# =====================================================================
# 1. TRANSACTIONS
# =====================================================================
sheet_title(tx, "Transactions",
            "The only tab where you type. Add new rows at the FIRST EMPTY ROW at the bottom - "
            "do not insert rows in the middle. To put things in date order, click the arrow on "
            "Date and sort; the Cash balance recalculates itself either way.", 10)

TX_HDRS = ["Date", "Description", "Category", "Money In", "Money Out",
           "Account", "Loan / Savings bucket", "Notes", "Month", "Cash balance",
           "k", "l", "u", "net"]
TX_W = [12, 32, 19, 12, 12, 13, 21, 26, 11, 14, 6, 6, 6, 6]
table_head(tx, 4, 1, TX_HDRS, TX_W, height=32)
for col in (9, 10):
    tx.cell(row=4, column=col).fill = PatternFill("solid", fgColor="6E7C8A")
tx.freeze_panes = "A5"

# Your real rows, carried over. "Bill" corrected to "Bills" and "Credit card"
# to "Credit Card" so they match the dropdown lists exactly.
examples = [
    (dt.datetime(2026, 8, 15), "World Bank paycheck", "Income", 9867.73, None, "Checking", None, "23rd July to 15th Aug, +5k relocation bonus"),
    (dt.datetime(2026, 8, 15), "Credit card repayment", "Bills", None, 3208.29, "Credit Card", None, "Chase"),
    (dt.datetime(2026, 8, 15), "Climb Launch payment", "Student Debt", None, 643.87, "Checking", "Climb Launch", None),
    (dt.datetime(2026, 8, 15), "Climb UAS payment", "Student Debt", None, 568.37, "Checking", "Climb UAS", None),
    (dt.datetime(2026, 8, 15), "Brokerage transfer", "Savings/Investments", None, 900, "Checking", "Brokerage / Stocks", None),
    (dt.datetime(2026, 8, 15), "Masters fund", "Savings/Investments", None, 600, "Savings", "Masters Fund", None),
    (dt.datetime(2026, 8, 16), "Groceries", "Food", None, 85.42, "Checking", None, None),
]

for r in range(TX0, TX1 + 1):
    i = r - TX0
    if i < len(examples):
        d, desc, cat, mi, mout, acct, bucket, note = examples[i]
        tx.cell(row=r, column=1, value=d)
        tx.cell(row=r, column=2, value=desc)
        tx.cell(row=r, column=3, value=cat)
        if mi is not None:
            tx.cell(row=r, column=4, value=mi)
        if mout is not None:
            tx.cell(row=r, column=5, value=mout)
        tx.cell(row=r, column=6, value=acct)
        if bucket:
            tx.cell(row=r, column=7, value=bucket)
        tx.cell(row=r, column=8, value=note)

    for col in range(1, 9):
        c = tx.cell(row=r, column=col)
        c.fill = FILL_INPUT
        c.border = BOX
        c.font = F_BODY
    tx.cell(row=r, column=1).number_format = DATEF
    tx.cell(row=r, column=4).number_format = MONEY
    tx.cell(row=r, column=5).number_format = MONEY

    cell(tx, r, 9, f'=IF($A{r}="","",EOMONTH($A{r},0))', MONF, fill=FILL_BAND, align="center")
    # Cash balance sums an absolute-anchored range rather than chaining off the
    # row above, so inserting, deleting, moving or sorting rows cannot break it.
    cell(tx, r, 10, f'=IF(COUNTA($A{r}:$H{r})=0,"",SUM($N${TX0}:$N{r}))',
         MONEY, font=F_BOLD, fill=FILL_BAND)
    tx.cell(row=r, column=11, value=f'=IF($A{r}="",0,$A{r}*100000+ROW())')
    tx.cell(row=r, column=12, value=(
        f'=IF(AND($C{r}="Student Debt",$G{r}="Climb Launch"),'
        f'COUNTIFS($C${TX0}:$C${TX1},"Student Debt",$G${TX0}:$G${TX1},"Climb Launch",'
        f'$K${TX0}:$K${TX1},"<="&$K{r}),"")'))
    tx.cell(row=r, column=13, value=(
        f'=IF(AND($C{r}="Student Debt",$G{r}="Climb UAS"),'
        f'COUNTIFS($C${TX0}:$C${TX1},"Student Debt",$G${TX0}:$G${TX1},"Climb UAS",'
        f'$K${TX0}:$K${TX1},"<="&$K{r}),"")'))
    # net effect of this row on total cash; a Transfer moves money between your
    # own accounts, so it nets to zero and never moves the balance
    tx.cell(row=r, column=14, value=(
        f'=IF($C{r}="Transfer",0,IF($D{r}="",0,$D{r})-IF($E{r}="",0,$E{r}))')).number_format = MONEY

for col in ("K", "L", "M", "N"):
    tx.column_dimensions[col].hidden = True

tx.conditional_formatting.add(f"A{TX0}:H{TX1}", FormulaRule(
    formula=[f'AND($C{TX0}="Student Debt",$G{TX0}="")'],
    fill=PatternFill("solid", bgColor=BADF), stopIfTrue=False))
tx.conditional_formatting.add(f"A{TX0}:H{TX1}", FormulaRule(
    formula=[f'AND($C{TX0}="Savings/Investments",$G{TX0}="")'],
    fill=PatternFill("solid", bgColor=WARNF), stopIfTrue=False))
tx.conditional_formatting.add(f"A{TX0}:H{TX1}", FormulaRule(
    formula=[f'AND($D{TX0}<>"",$E{TX0}<>"")'],
    fill=PatternFill("solid", bgColor=WARNF), stopIfTrue=False))
tx.auto_filter.ref = f"A4:J{TX1}"

# =====================================================================
# 6. START HERE  (guide + the one-time numbers)
# =====================================================================
sheet_title(sh, "Start Here",
            "How this workbook works, and the handful of numbers that were set up once.", 8)
for col, w in zip("ABCDEFGHIJ", [46, 16, 17, 17, 11, 14, 4, 24, 24, 24]):
    sh.column_dimensions[col].width = w

banner(sh, 4, 1, 6, "HOW TO USE THIS")
guide = [
    ("1.", "Everything you do happens on the Transactions tab. One row per transaction."),
    ("", "Date, Description, Category (dropdown), then Money In or Money Out. That's it."),
    ("2.", "Paying a loan? Category = Student Debt, and pick the loan in 'Loan / Savings bucket'."),
    ("3.", "Saving money? Category = Savings/Investments, and pick where it went in that same column."),
    ("", "That's how the Savings tab knows how much sits in stocks vs your masters fund."),
    ("4.", "Every other tab calculates itself. You never type on them."),
    ("", "Dashboard = the big picture.  Loan Payoff = what to pay and what you owe."),
    ("", "Monthly = month by month + spending by category.  Savings = where your savings go."),
    ("5.", "Tabs with a yellow box have a dropdown you can change (month pickers, savings split)."),
]
for i, (n, t) in enumerate(guide):
    r = 5 + i
    c = sh.cell(row=r, column=1, value=(n + " " + t) if n else ("     " + t))
    c.font = F_BOLD if n else Font(size=10, color=GREY)
    sh.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    sh.row_dimensions[r].height = 15

banner(sh, 15, 1, 6, "YOUR LOANS  (set once, refresh from a statement now and then)")
table_head(sh, 16, 1, ["Loan", "Principal", "Accrued interest", "Opening balance",
                       "APR", "Balance as of"], None, height=30)
loan_seed = [("Climb Launch", 14485.33, 47.89, 0.0625),
             ("Climb UAS", 12790.67, 38.56, 0.0625)]
for i, (nm, prin, acc, apr) in enumerate(loan_seed):
    r = 17 + i
    cell(sh, r, 1, nm, font=F_BOLD, fill=FILL_BAND)
    cell(sh, r, 2, prin, MONEY, fill=FILL_INPUT)
    cell(sh, r, 3, acc, MONEY, fill=FILL_INPUT)
    cell(sh, r, 4, f"=$B{r}+$C{r}", MONEY, font=F_BOLD, fill=FILL_CARD)
    cell(sh, r, 5, apr, PCT2, fill=FILL_INPUT, align="center")
    cell(sh, r, 6, START, DATEF, fill=FILL_INPUT, align="center")
cell(sh, 19, 1, "TOTAL", font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND)
for col in (2, 3, 4):
    L = get_column_letter(col)
    cell(sh, 19, col, f"=SUM({L}17:{L}18)", MONEY,
         font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND)
cell(sh, 19, 5, "", fill=FILL_BAND)
cell(sh, 19, 6, "", fill=FILL_BAND)

# NOTE: OPEN_L/OPEN_U/OPEN_T point at D17/D18/D19
OPEN_L, OPEN_U, OPEN_T = f"{SH}!$D$17", f"{SH}!$D$18", f"{SH}!$D$19"
APR_L, APR_U = f"{SH}!$E$17", f"{SH}!$E$18"
ASOF_L, ASOF_U = f"{SH}!$F$17", f"{SH}!$F$18"

banner(sh, 21, 1, 6, "THE GOAL")
cell(sh, 22, 1, "Be debt-free by", font=F_BOLD, fill=FILL_BAND)
cell(sh, 22, 2, TARGET, DATEF, font=Font(bold=True, size=11, color=INK),
     fill=FILL_INPUT, align="center")
cell(sh, 23, 1, "Payments still to make", font=F_BOLD, fill=FILL_BAND)
cell(sh, 23, 2, (
    f'=MAX(1,(YEAR($B$22)-YEAR(TODAY()))*12+(MONTH($B$22)-MONTH(TODAY()))'
    f'+IF(SUMIFS({TXOUT},{TXD},">="&EOMONTH(TODAY(),-1)+1,{TXD},"<="&EOMONTH(TODAY(),0),'
    f'{TXC},"Student Debt")>0,0,1))'), INT,
     font=Font(bold=True, size=11, color=INK), fill=FILL_CARD, align="center")
TARGET_DATE, MONTHS_LEFT = f"{SH}!$B$22", f"{SH}!$B$23"
cell(sh, 24, 1, "The current month counts only if you have not paid it yet.",
     font=F_NOTE, border=NOBOX)

banner(sh, 26, 1, 6, "MONTHLY BUDGET PLAN  (what you intend to do)")
table_head(sh, 27, 1, ["Line", "Planned monthly", "Category it maps to"], None, height=22)
budget = [("Take-home income", 6000, "Income"),
          ("Rent", 1800, "Housing"),
          ("Phone + gym + internet", 200, "Bills"),
          ("Transportation", 120, "Transportation"),
          ("Food / groceries", 0, "Food"),
          ("Family", 1250, "Family"),
          ("Savings (all buckets)", 1500, "Savings/Investments"),
          ("Student loans", f"={REQ_T}", "Student Debt")]
for i, (lab, v, cat) in enumerate(budget):
    r = 28 + i
    cell(sh, r, 1, lab, font=F_BOLD if i == 0 else F_BODY, fill=FILL_BAND)
    cell(sh, r, 2, v, MONEY,
         fill=FILL_CARD if isinstance(v, str) else FILL_INPUT,
         font=F_BOLD if isinstance(v, str) else F_BODY)
    cell(sh, r, 3, cat, font=F_NOTE)
cell(sh, 36, 1, "Total planned out", font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND)
cell(sh, 36, 2, "=SUM($B$29:$B$35)", MONEY,
     font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND)
cell(sh, 37, 1, "Left over / (short)", font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND)
cell(sh, 37, 2, "=$B$28-$B$36", MONEY,
     font=Font(bold=True, size=11, color=INK), fill=FILL_BAND)
sh.conditional_formatting.add("B37", FormulaRule(
    formula=['$B$37<0'], fill=PatternFill("solid", bgColor=BADF), font=Font(bold=True, color=RED)))
sh.conditional_formatting.add("B37", FormulaRule(
    formula=['$B$37>=0'], fill=PatternFill("solid", bgColor=GOODF), font=Font(bold=True, color=GREEN)))
cell(sh, 38, 1, "Food is $0 because your original plan had no groceries line. "
                "Put a real number in and watch 'Left over' react.", font=F_NOTE, border=NOBOX)
PLAN_INCOME, PLAN_SAVE, PLAN_LOAN = f"{SH}!$B$28", f"{SH}!$B$34", f"{SH}!$B$35"
PLAN_OUT, PLAN_LEFT = f"{SH}!$B$36", f"{SH}!$B$37"

banner(sh, 40, 1, 6, "DATA HEALTH  (all zeros = your Transactions tab is clean)")
checks = [
    ("Loan payments with no loan picked  (red rows)",
     f'=COUNTIFS({TXC},"Student Debt",{TXG},"")'),
    ("Savings with no bucket picked  (amber rows)",
     f'=COUNTIFS({TXC},"Savings/Investments",{TXG},"")'),
    ("Rows with both Money In and Money Out",
     f'=SUMPRODUCT(({TXIN}<>"")*({TXOUT}<>""))'),
    ("Rows with an amount but no category",
     f'=SUMPRODUCT((({TXIN}<>"")+({TXOUT}<>"")>0)*({TXC}=""))'),
]
for i, (lab, f) in enumerate(checks):
    r = 41 + i
    cell(sh, r, 1, lab, fill=FILL_BAND)
    cell(sh, r, 2, f, INT, font=F_BOLD, fill=FILL_CARD, align="center")
sh.conditional_formatting.add("B41:B44", FormulaRule(
    formula=['B41>0'], fill=PatternFill("solid", bgColor=BADF), font=Font(bold=True, color=RED)))
sh.conditional_formatting.add("B41:B44", FormulaRule(
    formula=['B41=0'], fill=PatternFill("solid", bgColor=GOODF), font=Font(bold=True, color=GREEN)))

banner(sh, 46, 1, 6, "STATEMENT CHECK  (optional - keeps the tracker honest)")
table_head(sh, 47, 1, ["Loan", "Servicer balance", "As of", "Tracker says", "Difference"],
           None, height=22)
for i, (nm, ref) in enumerate([("Climb Launch", BAL_L), ("Climb UAS", BAL_U)]):
    r = 48 + i
    cell(sh, r, 1, nm, font=F_BOLD, fill=FILL_BAND)
    cell(sh, r, 2, None, MONEY, fill=FILL_INPUT)
    cell(sh, r, 3, None, DATEF, fill=FILL_INPUT, align="center")
    cell(sh, r, 4, f"={ref}", MONEY, fill=FILL_CARD)
    cell(sh, r, 5, f'=IF($B{r}="","",$B{r}-$D{r})', MONEY, font=F_BOLD, fill=FILL_CARD)
cell(sh, 50, 1, "Private loans accrue daily so a few dollars of drift is normal. "
                "If the gap grows, refresh Principal and Accrued interest above. "
                "Note: your Climb UAS statement figures do not quite reconcile - "
                "12,790.67 + 38.56 = 12,829.23 but you quoted 12,833.61.",
     font=F_NOTE, border=NOBOX)
sh.merge_cells("A50:F52")
sh["A50"].alignment = Alignment(vertical="top", wrap_text=True)

# dropdown source lists, editable. Each list gets SPARE rows on the end so you
# can add your own option and the dropdown picks it up with no other changes.
SPARE = 5
banner(sh, 4, 8, 10, "DROPDOWN LISTS")
table_head(sh, 5, 8, ["Category", "Account", "Loan / Savings bucket"], None, height=22)
LIST0 = 6
LIST1 = LIST0 + max(len(CATEGORIES), len(ACCOUNTS), len(BUCKETS)) + SPARE - 1
for i in range(LIST1 - LIST0 + 1):
    r = LIST0 + i
    for col, src in ((8, CATEGORIES), (9, ACCOUNTS), (10, BUCKETS)):
        cell(sh, r, col, src[i] if i < len(src) else None, fill=FILL_INPUT)
cell(sh, LIST1 + 1, 8,
     "Type into a blank cell above to add your own option.",
     font=F_NOTE, border=NOBOX)

for nm, ref in (("CatList", f"{SH}!$H${LIST0}:$H${LIST0+len(CATEGORIES)+SPARE-1}"),
                ("AcctList", f"{SH}!$I${LIST0}:$I${LIST0+len(ACCOUNTS)+SPARE-1}"),
                ("BucketList", f"{SH}!$J${LIST0}:$J${LIST0+len(BUCKETS)+SPARE-1}")):
    wb.defined_names.add(DefinedName(nm, attr_text=ref))

# The lists are mirrored into hidden columns on THIS sheet and the dropdowns point
# at the mirror. A same-sheet range is the most portable form of list validation -
# a cross-sheet one gets shoved into an Excel-only extension block.
# showErrorMessage=True is the bit that makes the dropdown feel connected: without
# it Excel shows the arrow but still accepts anything you type over the top, which
# is how "Bill" got in instead of "Bills".
MIRROR = [(16, "P", CATEGORIES, 8), (17, "Q", ACCOUNTS, 9), (18, "R", BUCKETS, 10)]
for col, L, src, sh_col in MIRROR:
    for i in range(len(src) + SPARE):
        r = TX0 + i
        sr = LIST0 + i
        tx.cell(row=r, column=col,
                value=f'=IF({SH}!${get_column_letter(sh_col)}${sr}="","",'
                      f'{SH}!${get_column_letter(sh_col)}${sr})')
    tx.column_dimensions[L].hidden = True

for (col, L, src, _), tgt in zip(MIRROR, (f"C{TX0}:C{TX1}", f"F{TX0}:F{TX1}", f"G{TX0}:G{TX1}")):
    dv = DataValidation(type="list", allow_blank=True, showDropDown=False,
                        formula1=f"${L}${TX0}:${L}${TX0+len(src)+SPARE-1}",
                        showErrorMessage=True, errorStyle="stop")
    dv.errorTitle = "Use the dropdown"
    dv.error = ("That is not one of your options, so the rest of the workbook would "
                "not count it. Click the arrow in this cell and pick from the list. "
                "To create a new option, add it on the Start Here tab.")
    dv.promptTitle = "Pick from the list"
    dv.prompt = "Click the arrow on the right of this cell."
    tx.add_data_validation(dv)
    dv.add(tgt)

# =====================================================================
# 3. LOAN PAYOFF  (the plan + every payment, one tab)
# =====================================================================
sheet_title(lp, "Loan Payoff",
            "Top: what to pay. Middle: month by month. Bottom: every single payment with the "
            "balance after it. All of it comes from the Transactions tab.", 13)
LPW = [11, 13, 24, 13, 14, 13, 14, 13, 13, 13, 14, 13, 16]
for i, w in enumerate(LPW):
    lp.column_dimensions[get_column_letter(i + 1)].width = w

MT0, MT1 = 12, 12 + N_MONTHS - 1                 # month table rows 12..47
LD_L_OPEN, LD_L0 = 51, 52
LD_L1 = LD_L0 + N_PAY - 1                        # 52..91
LD_U_OPEN, LD_U0 = 95, 96
LD_U1 = LD_U0 + N_PAY - 1                        # 96..135

BAL_L, REQ_L = f"{LP}!$A$6", f"{LP}!$B$6"
BAL_U, REQ_U = f"{LP}!$C$6", f"{LP}!$D$6"
BAL_T, REQ_T = f"{LP}!$E$6", f"{LP}!$F$6"
TGT_D, PROJ_D = f"{LP}!$G$6", f"{LP}!$H$6"
PAID_L, INT_L, NPAY_L = f"{LP}!$M$51", f"{LP}!$M$52", f"{LP}!$M$53"
PAID_U, INT_U, NPAY_U = f"{LP}!$M$95", f"{LP}!$M$96", f"{LP}!$M$97"
PAID_T = f"({PAID_L}+{PAID_U})"
INT_T = f"({INT_L}+{INT_U})"
PCT_DONE = f"IF({OPEN_T}=0,0,({OPEN_T}-{BAL_T})/{OPEN_T})"

banner(lp, 4, 1, 13, "THE PLAN")
cards = [
    ("Climb Launch owes", f"=LOOKUP(9.99E+307,$H${LD_L_OPEN}:$H${LD_L1})", MONEY0),
    ("pay / month", f"=IFERROR(-PMT({APR_L}/12,{MONTHS_LEFT},$A$6),0)", MONEY),
    ("Climb UAS owes", f"=LOOKUP(9.99E+307,$H${LD_U_OPEN}:$H${LD_U1})", MONEY0),
    ("pay / month", f"=IFERROR(-PMT({APR_U}/12,{MONTHS_LEFT},$C$6),0)", MONEY),
    ("TOTAL owed", "=$A$6+$C$6", MONEY0),
    ("TOTAL / month", "=$B$6+$D$6", MONEY),
    ("debt-free target", f"={TARGET_DATE}", MONF),
    ("projected finish", (
        f'=IFERROR(INDEX($A${MT0}:$A${MT1},'
        f'MATCH(TRUE,INDEX($K${MT0}:$K${MT1}<=0.005,0),0)),"")'), MONF),
]
for i, (lab, f, fmt) in enumerate(cards):
    kpi(lp, 5, i + 1, lab, f, fmt, small=(i in (1, 3, 5)))
for col in (2, 4, 6):
    lp.cell(row=6, column=col).font = Font(bold=True, size=13, color=RED)
lp.cell(row=6, column=7).font = Font(bold=True, size=12, color=NAVY)
lp.cell(row=6, column=8).font = Font(bold=True, size=12, color=GREEN)

lp.merge_cells("A7:H7")
c = lp.cell(row=7, column=1)
c.value = (f'="Pay $"&TEXT($F$6,"#,##0")&" a month ($"&TEXT($B$6,"#,##0")&" to Launch, $"'
           f'&TEXT($D$6,"#,##0")&" to UAS) for the next "&{MONTHS_LEFT}&" months. "'
           f'&IF($H$6="","At your current pace this runs past the chart - pay more.",'
           f'IF($H$6<=EOMONTH($G$6,0),"You are ON TRACK for "&TEXT($G$6,"mmm yyyy")&".",'
           f'"At your current pace you finish "&TEXT($H$6,"mmm yyyy")&" - that is "'
           f'&((YEAR($H$6)-YEAR($G$6))*12+MONTH($H$6)-MONTH($G$6))&" months LATE."))')
c.font = Font(bold=True, size=11, color=INK)
c.alignment = Alignment(vertical="center", indent=1)
lp.row_dimensions[7].height = 22
lp.conditional_formatting.add("A7:H7", FormulaRule(
    formula=['AND($H$6<>"",$H$6<=EOMONTH($G$6,0))'],
    fill=PatternFill("solid", bgColor=GOODF), stopIfTrue=True))
lp.conditional_formatting.add("A7:H7", FormulaRule(
    formula=['TRUE'], fill=PatternFill("solid", bgColor=BADF)))

cell(lp, 8, 1, "PAID OFF SO FAR", font=F_LBL, fill=FILL_CARD, align="center")
bar_cell(lp, 8, 2, 8, PCT_DONE, color=GREEN, n=42)
kpi(lp, 5, 10, "paid to date", "=$M$51+$M$95", MONEY0, small=True, span=2)
kpi(lp, 5, 12, "interest accrued", "=$M$52+$M$96", MONEY0, small=True, span=2)

# ---- month by month
banner(lp, 10, 1, 13, "MONTH BY MONTH")
table_head(lp, 11, 1, ["Month", "Target\nLaunch", "Target\nUAS", "Target\nTOTAL",
                       "Paid\nLaunch", "Paid\nUAS", "Paid\nTOTAL", "Variance",
                       "Launch\nbalance", "UAS\nbalance", "TOTAL\nbalance",
                       "If always\non target", "Status"], None, height=32)
lp.freeze_panes = "A12"

for i in range(N_MONTHS):
    r = MT0 + i
    cell(lp, r, 1, START if i == 0 else f"=EDATE($A{r-1},1)", MONF, font=F_BOLD, align="center")
    if i == 0:
        cell(lp, r, 2, "=$B$6", MONEY, fill=FILL_BAND)
        cell(lp, r, 3, "=$D$6", MONEY, fill=FILL_BAND)
    else:
        cell(lp, r, 2, f'=IF($P{r-1}<=0.005,0,$B$6)', MONEY, fill=FILL_BAND)
        cell(lp, r, 3, f'=IF($Q{r-1}<=0.005,0,$D$6)', MONEY, fill=FILL_BAND)
    cell(lp, r, 4, f"=$B{r}+$C{r}", MONEY, font=F_BOLD, fill=FILL_BAND)
    for col, bucket in ((5, "Climb Launch"), (6, "Climb UAS")):
        crit = f',{TXC},"Student Debt",{TXG},"{bucket}"'
        cell(lp, r, col, "=" + m_out(f"$A{r}", crit), MONEY)
    cell(lp, r, 7, f"=$E{r}+$F{r}", MONEY, font=F_BOLD)
    cell(lp, r, 8, f"=$G{r}-$D{r}", MONEY)

    fbL = "$B$6" if i == 0 else f'IF($I{r-1}<=0.005,0,$B$6)'
    fbU = "$D$6" if i == 0 else f'IF($J{r-1}<=0.005,0,$D$6)'
    lp.cell(row=r, column=14, value=(
        f'=IF($A{r}<=EOMONTH(TODAY(),-1),$E{r},IF($E{r}>0,$E{r},{fbL}))')).number_format = MONEY
    lp.cell(row=r, column=15, value=(
        f'=IF($A{r}<=EOMONTH(TODAY(),-1),$F{r},IF($F{r}>0,$F{r},{fbU}))')).number_format = MONEY

    this_month = f'$A{r}=DATE(YEAR(TODAY()),MONTH(TODAY()),1)'
    if i == 0:
        recL = f'MAX(0,{OPEN_L}*(1+{APR_L}/12)-$N{r})'
        recU = f'MAX(0,{OPEN_U}*(1+{APR_U}/12)-$O{r})'
        lp.cell(row=r, column=16, value=f'=MAX(0,{OPEN_L}*(1+{APR_L}/12)-$B{r})').number_format = MONEY
        lp.cell(row=r, column=17, value=f'=MAX(0,{OPEN_U}*(1+{APR_U}/12)-$C{r})').number_format = MONEY
    else:
        recL = f'IF($I{r-1}<=0.005,0,MAX(0,$I{r-1}*(1+{APR_L}/12)-$N{r}))'
        recU = f'IF($J{r-1}<=0.005,0,MAX(0,$J{r-1}*(1+{APR_U}/12)-$O{r}))'
        lp.cell(row=r, column=16, value=f'=IF($P{r-1}<=0.005,0,MAX(0,$P{r-1}*(1+{APR_L}/12)-$B{r}))').number_format = MONEY
        lp.cell(row=r, column=17, value=f'=IF($Q{r-1}<=0.005,0,MAX(0,$Q{r-1}*(1+{APR_U}/12)-$C{r}))').number_format = MONEY
    cell(lp, r, 9, f"=IF({this_month},$A$6,{recL})", MONEY)
    cell(lp, r, 10, f"=IF({this_month},$C$6,{recU})", MONEY)
    cell(lp, r, 11, f"=$I{r}+$J{r}", MONEY, font=F_BOLD)
    cell(lp, r, 12, f"=$P{r}+$Q{r}", MONEY, font=Font(size=10, color=GREY))
    cell(lp, r, 13, (
        f'=IF($K{r}<=0.005,"DEBT FREE",'
        f'IF($A{r}>EOMONTH(TODAY(),0),"projected",'
        f'IF($H{r}>=-1,"on track","behind "&TEXT(-$H{r},"$#,##0"))))'), align="center")

for col in ("N", "O", "P", "Q"):
    lp.column_dimensions[col].hidden = True

lp.conditional_formatting.add(f"M{MT0}:M{MT1}", FormulaRule(
    formula=[f'LEFT($M{MT0},6)="behind"'],
    fill=PatternFill("solid", bgColor=BADF), font=Font(bold=True, color=RED)))
lp.conditional_formatting.add(f"M{MT0}:M{MT1}", FormulaRule(
    formula=[f'$M{MT0}="on track"'],
    fill=PatternFill("solid", bgColor=GOODF), font=Font(bold=True, color=GREEN)))
lp.conditional_formatting.add(f"M{MT0}:M{MT1}", FormulaRule(
    formula=[f'$M{MT0}="DEBT FREE"'],
    fill=PatternFill("solid", bgColor="DCEAF7"), font=Font(bold=True, color=NAVY)))
lp.conditional_formatting.add(f"K{MT0}:K{MT1}", DataBarRule(
    start_type="num", start_value=0, end_type="num", end_value=28000,
    color="C4756B", showValue=True))

# ---- per-payment ledgers
LED_HDRS = ["#", "Date", "Description", "Days", "Interest\naccrued",
            "Payment", "To\nprincipal", "Balance after"]
for nm, bnr, opn, r0, r1, apr, openref, asof, seqcol, sumrow in (
        ("Climb Launch", 49, LD_L_OPEN, LD_L0, LD_L1, APR_L, OPEN_L, ASOF_L, "L", 51),
        ("Climb UAS", 93, LD_U_OPEN, LD_U0, LD_U1, APR_U, OPEN_U, ASOF_U, "M", 95)):
    banner(lp, bnr, 1, 8, f"EVERY PAYMENT  -  {nm.upper()}")
    table_head(lp, bnr + 1, 1, LED_HDRS, None, height=30)
    cell(lp, opn, 1, 0, INT, fill=FILL_BAND, align="center")
    cell(lp, opn, 2, f"={asof}", DATEF, fill=FILL_BAND, align="center")
    cell(lp, opn, 3, "Opening balance", font=F_BOLD, fill=FILL_BAND)
    for col in (4, 5, 6, 7):
        cell(lp, opn, col, None, fill=FILL_BAND)
    cell(lp, opn, 8, f"={openref}", MONEY, font=F_BOLD, fill=FILL_BAND)

    for r in range(r0, r1 + 1):
        seq = r - opn
        m = f'MATCH({seq},{TXS}!${seqcol}${TX0}:${seqcol}${TX1},0)'
        cell(lp, r, 1, f'=IF($B{r}="","",{seq})', INT, align="center")
        cell(lp, r, 2, f'=IFERROR(INDEX({TXS}!$A${TX0}:$A${TX1},{m}),"")', DATEF, align="center")
        cell(lp, r, 3, f'=IFERROR(INDEX({TXS}!$B${TX0}:$B${TX1},{m}),"")')
        cell(lp, r, 4, f'=IF($B{r}="","",$B{r}-$B{r-1})', INT, align="center")
        cell(lp, r, 5, f'=IF($B{r}="","",$H{r-1}*{apr}/365*$D{r})', MONEY,
             font=Font(size=10, color=AMBER))
        cell(lp, r, 6, f'=IFERROR(INDEX({TXS}!$E${TX0}:$E${TX1},{m}),"")', MONEY)
        cell(lp, r, 7, f'=IF($B{r}="","",$F{r}-$E{r})', MONEY,
             font=Font(size=10, color=GREEN))
        cell(lp, r, 8, f'=IF($B{r}="","",MAX(0,$H{r-1}+$E{r}-$F{r}))', MONEY, font=F_BOLD)
    zebra(lp, f"A{r0}:H{r1}", r0)
    lp.conditional_formatting.add(f"H{r0}:H{r1}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=15000,
        color="7BA7CC", showValue=True))

    for j, (lab, f) in enumerate((("Total paid", f"=SUM($F${r0}:$F${r1})"),
                                  ("Interest accrued", f"=SUM($E${r0}:$E${r1})"),
                                  ("Payments logged", f"=COUNT($F${r0}:$F${r1})"))):
        cell(lp, sumrow + j, 12, lab, font=F_LBL, fill=FILL_CARD, align="right")
        cell(lp, sumrow + j, 13, f, INT if j == 2 else MONEY, font=F_BOLD,
             fill=FILL_CARD, align="center")

cell(lp, LD_U1 + 2, 1,
     "Interest is worked out daily: balance x APR / 365 x days since the previous payment. "
     "That is how private loans actually work, so this matches your statement closely.",
     font=F_NOTE, border=NOBOX)

# =====================================================================
# 2. DASHBOARD
# =====================================================================
sheet_title(db, "Dashboard",
            "Nothing to fill in here. Change the month in the yellow box to look at a "
            "different month.", 10)
for i, w in enumerate([20, 17, 17, 17, 17, 17, 17, 17, 3, 3]):
    db.column_dimensions[get_column_letter(i + 1)].width = w

DSEL = month_picker(db, 3, 2, 20)          # month list hidden in column T
DSEL = f"$B$3"
cell(db, 3, 4, '="Showing "&TEXT($B$3,"mmmm yyyy")&"   |   today is "&TEXT(TODAY(),"mmm d, yyyy")',
     font=Font(italic=True, size=10, color=GREY), border=NOBOX)

# --- debt block
banner(db, 5, 1, 8, "DEBT")
kpi(db, 6, 1, "still owed", f"={BAL_T}", MONEY0)
kpi(db, 6, 2, "paid off", "=" + PCT_DONE, PCT)
kpi(db, 6, 3, "paid to date", f"={PAID_L}+{PAID_U}", MONEY0)
kpi(db, 6, 4, "interest accrued", f"={INT_L}+{INT_U}", MONEY0)
kpi(db, 6, 5, "interest / month", f"=({BAL_L}*{APR_L}+{BAL_U}*{APR_U})/12", MONEY0)
kpi(db, 6, 6, "climb launch", f"={BAL_L}", MONEY0, small=True)
kpi(db, 6, 7, "climb uas", f"={BAL_U}", MONEY0, small=True)
kpi(db, 6, 8, "started at", f"={OPEN_T}", MONEY0, small=True)
db.cell(row=7, column=1).font = Font(bold=True, size=16, color=RED)
db.cell(row=7, column=2).font = Font(bold=True, size=16, color=GREEN)
cell(db, 9, 1, "PROGRESS", font=F_LBL, fill=FILL_CARD, align="center")
bar_cell(db, 9, 2, 8, PCT_DONE, color=GREEN, n=44)

# --- on track block
banner(db, 11, 1, 8, "ARE YOU GOING TO MAKE JULY 2028?")
kpi(db, 12, 1, "target", f"={TARGET_DATE}", MONF, small=True)
kpi(db, 12, 2, "payments left", f"={MONTHS_LEFT}", INT, small=True)
kpi(db, 12, 3, "required / month", f"={REQ_T}", MONEY, small=True)
kpi(db, 12, 4, "projected finish", f'=IF({PROJ_D}="","past 2029",{PROJ_D})', MONF, small=True)
v = kpi(db, 12, 5, "verdict", (
    f'=IF({PROJ_D}="","BEHIND",IF({PROJ_D}<=EOMONTH({TARGET_DATE},0),"ON TRACK",'
    f'"LATE BY "&((YEAR({PROJ_D})-YEAR({TARGET_DATE}))*12'
    f'+MONTH({PROJ_D})-MONTH({TARGET_DATE}))&" MO"))'), TEXTF, small=True, span=4)
v.alignment = Alignment(horizontal="center", vertical="center")
db.conditional_formatting.add("E13:H13", FormulaRule(
    formula=['$E$13="ON TRACK"'], fill=PatternFill("solid", bgColor=GOODF),
    font=Font(bold=True, size=13, color=GREEN)))
db.conditional_formatting.add("E13:H13", FormulaRule(
    formula=['$E$13<>"ON TRACK"'], fill=PatternFill("solid", bgColor=BADF),
    font=Font(bold=True, size=13, color=RED)))

# --- this month's marching orders
banner(db, 15, 1, 8, "WHAT TO PAY IN THE MONTH SHOWN ABOVE")
table_head(db, 16, 1, ["Loan", "Target", "Paid so far", "Still to pay", "Balance now"],
           None, height=22)
for i, (nm, tgt, bal) in enumerate([("Climb Launch", REQ_L, BAL_L),
                                    ("Climb UAS", REQ_U, BAL_U)]):
    r = 17 + i
    crit = f',{TXC},"Student Debt",{TXG},"{nm}"'
    cell(db, r, 1, nm, font=F_BOLD, fill=FILL_BAND)
    cell(db, r, 2, f"={tgt}", MONEY, fill=FILL_BAND)
    cell(db, r, 3, "=" + m_out(DSEL, crit), MONEY)
    cell(db, r, 4, f"=MAX(0,$B{r}-$C{r})", MONEY, font=Font(bold=True, size=10, color=RED))
    cell(db, r, 5, f"={bal}", MONEY, font=F_BOLD)
cell(db, 19, 1, "TOTAL", font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND)
for col in range(2, 6):
    L = get_column_letter(col)
    cell(db, 19, col, f"=SUM({L}17:{L}18)", MONEY,
         font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND)

# --- cash for the selected month
banner(db, 21, 1, 8, "CASH IN THE MONTH SHOWN ABOVE")
table_head(db, 22, 1, ["Money in", "Living costs", "Loan payments", "Savings",
                       "Left over", "Plan: loans", "Plan: savings", "Plan: left over"],
           None, height=30)
sd = f',{TXC},"Student Debt"'
ss = f',{TXC},"Savings/Investments"'
stf = f',{TXC},"Transfer"'
cells23 = [
    "=" + m_in(DSEL),
    "=" + m_out(DSEL) + " -$C$23-$D$23",
    "=" + m_out(DSEL, sd),
    "=" + m_out(DSEL, ss),
    "=$A$23-$B$23-$C$23-$D$23",
    f"={REQ_T}",
    f"={PLAN_SAVE}",
    f"={PLAN_LEFT}",
]
for i, f in enumerate(cells23):
    cell(db, 23, i + 1, f, MONEY, font=Font(bold=True, size=11, color=INK),
         fill=FILL_CARD if i < 5 else FILL_BAND, align="center")
db.conditional_formatting.add("E23", FormulaRule(
    formula=['$E$23<0'], font=Font(bold=True, size=11, color=RED)))

# --- savings block
banner(db, 25, 1, 8, "SAVINGS")
kpi(db, 26, 1, "saved to date", f'=SUMIFS({TXOUT},{TXC},"Savings/Investments")', MONEY0)
kpi(db, 26, 2, "this month", "=$D$23", MONEY0)
kpi(db, 26, 3, "planned / month", f"={PLAN_SAVE}", MONEY0, small=True)
kpi(db, 26, 4, "biggest bucket", "=Savings!$A$21", TEXTF, small=True, span=2)
kpi(db, 26, 6, "in that bucket", "=Savings!$B$21", MONEY0, small=True)
kpi(db, 26, 7, "buckets used", "=Savings!$D$21", INT, small=True, span=2)
db.cell(row=27, column=1).font = Font(bold=True, size=16, color=GREEN)
cell(db, 28, 1, "See the Savings tab for the full split.", font=F_NOTE, border=NOBOX)

# --- charts
ch1 = LineChart()
ch1.title = "Loan balance - where you are vs. always paying the target"
ch1.style = 2
ch1.height, ch1.width = 8.5, 17
ch1.y_axis.numFmt = MONEY0
ch1.add_data(Reference(lp, min_col=11, max_col=12, min_row=11, max_row=MT1),
             titles_from_data=True)
ch1.set_categories(Reference(lp, min_col=1, min_row=MT0, max_row=MT1))
ch1.series[0].graphicalProperties.line.width = 28000
db.add_chart(ch1, "A30")

ch2 = BarChart()
ch2.type = "col"
ch2.grouping = "clustered"
ch2.title = "Money in vs living costs vs loans vs savings, by month"
ch2.style = 2
ch2.height, ch2.width = 8.5, 17
ch2.y_axis.numFmt = MONEY0
ch2.add_data(Reference(mo, min_col=2, max_col=5, min_row=5, max_row=41),
             titles_from_data=True)
ch2.set_categories(Reference(mo, min_col=1, min_row=6, max_row=41))
db.add_chart(ch2, "A48")

# =====================================================================
# 4. MONTHLY
# =====================================================================
sheet_title(mo, "Monthly",
            "Month by month, then spending by category for whichever month you pick, "
            "then your pay periods.", 11)
for i, w in enumerate([11, 13, 13, 14, 13, 13, 15, 13, 12, 14, 12]):
    mo.column_dimensions[get_column_letter(i + 1)].width = w

banner(mo, 4, 1, 11, "MONTH BY MONTH")
table_head(mo, 5, 1, ["Month", "Money in", "Living costs", "Loan payments", "Savings",
                      "Left over", "Cash balance", "Plan: loans", "vs plan",
                      "Plan: savings", "vs plan"], None, height=30)
mo.freeze_panes = "B6"
MM0, MM1 = 6, 6 + N_MONTHS - 1
for i in range(N_MONTHS):
    r = MM0 + i
    cell(mo, r, 1, START if i == 0 else f"=EDATE($A{r-1},1)", MONF, font=F_BOLD, align="center")
    sel = f"$A{r}"
    cell(mo, r, 2, "=" + m_in(sel), MONEY)
    cell(mo, r, 3, "=" + m_out(sel) + f" -$D{r}-$E{r}", MONEY)
    cell(mo, r, 4, "=" + m_out(sel, sd), MONEY)
    cell(mo, r, 5, "=" + m_out(sel, ss), MONEY)
    cell(mo, r, 6, f"=$B{r}-$C{r}-$D{r}-$E{r}", MONEY, font=F_BOLD)
    prev = "0" if i == 0 else f"$G{r-1}"
    cell(mo, r, 7, (
        f'=IFERROR(LOOKUP(2,1/(({TXD}>=$A{r})*({TXD}<EDATE($A{r},1))),'
        f'{TXS}!$J${TX0}:$J${TX1}),{prev})'), MONEY)
    cell(mo, r, 8, f"={REQ_T}", MONEY, fill=FILL_BAND)
    cell(mo, r, 9, f'=IF($D{r}=0,"",$D{r}-$H{r})', MONEY)
    cell(mo, r, 10, f"={PLAN_SAVE}", MONEY, fill=FILL_BAND)
    cell(mo, r, 11, f'=IF($E{r}=0,"",$E{r}-$J{r})', MONEY)

mo.conditional_formatting.add(f"F{MM0}:F{MM1}", FormulaRule(
    formula=[f'$F{MM0}<0'], font=Font(bold=True, color=RED)))
for rng in (f"I{MM0}:I{MM1}", f"K{MM0}:K{MM1}"):
    mo.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND(ISNUMBER({rng[0]}{MM0}),{rng[0]}{MM0}<-1)'],
        fill=PatternFill("solid", bgColor=BADF), font=Font(bold=True, color=RED)))
    mo.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND(ISNUMBER({rng[0]}{MM0}),{rng[0]}{MM0}>=-1)'],
        fill=PatternFill("solid", bgColor=GOODF), font=Font(bold=True, color=GREEN)))

# --- category spending with its own month picker
CB = MM1 + 2                                   # banner row
banner(mo, CB, 1, 7, "SPENDING BY CATEGORY")
MSEL = month_picker(mo, CB + 1, 2, 20)
MSEL = f"$B${CB+1}"
cell(mo, CB + 1, 4, '="Money out during "&TEXT(' + MSEL + ',"mmmm yyyy")',
     font=Font(italic=True, size=10, color=GREY), border=NOBOX)
table_head(mo, CB + 2, 1, ["Category", "Selected month", "Share of month",
                           "All time", "Monthly average"], None, height=22)
CR0 = CB + 3
CR1 = CR0 + len(SPEND_CATS) - 1
for i, cat in enumerate(SPEND_CATS):
    r = CR0 + i
    cell(mo, r, 1, cat, font=F_BOLD)
    cell(mo, r, 2, "=" + m_out(MSEL, f",{TXC},$A{r}"), MONEY)
    cell(mo, r, 3, f'=IFERROR($B{r}/SUM($B${CR0}:$B${CR1}),0)', PCT, align="center")
    cell(mo, r, 4, f'=SUMIFS({TXOUT},{TXC},$A{r})', MONEY)
    cell(mo, r, 5, (f'=IFERROR($D{r}/MAX(1,(YEAR(TODAY())-YEAR($A${MM0}))*12'
                    f'+MONTH(TODAY())-MONTH($A${MM0})+1),0)'), MONEY)
cell(mo, CR1 + 1, 1, "TOTAL", font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND)
for col in (2, 4):
    L = get_column_letter(col)
    cell(mo, CR1 + 1, col, f"=SUM({L}{CR0}:{L}{CR1})", MONEY,
         font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND)
cell(mo, CR1 + 1, 3, None, fill=FILL_BAND)
cell(mo, CR1 + 1, 5, None, fill=FILL_BAND)
zebra(mo, f"A{CR0}:E{CR1}", CR0)
mo.conditional_formatting.add(f"B{CR0}:B{CR1}", DataBarRule(
    start_type="num", start_value=0, end_type="max", color="6FA8C7", showValue=True))
mo.conditional_formatting.add(f"D{CR0}:D{CR1}", DataBarRule(
    start_type="num", start_value=0, end_type="max", color="C9C0DC", showValue=True))

# --- pay periods
PB = CR1 + 3
banner(mo, PB, 1, 9, "PAY PERIODS  (paid on the 15th and the 30th)")
table_head(mo, PB + 1, 1, ["Pay date", "Covers", "Money in", "Total out",
                           "Loans", "Savings", "Everything else", "Left over",
                           "Logged?"], None, height=22)

def semi_monthly(first, last_month):
    """(payday, period_end) pairs for pay on the 15th and the 30th.
    Short months pay on the last day instead; no day is ever double counted."""
    out, y, m = [], first.year, first.month
    while (y, m) <= last_month:
        eom = (dt.date(y + (m // 12), (m % 12) + 1, 1) - dt.timedelta(days=1)).day
        d2 = min(30, eom)                       # Feb pays on the 28th/29th
        nxt_y, nxt_m = y + (m // 12), (m % 12) + 1
        out.append((dt.datetime(y, m, 15), dt.datetime(y, m, d2) - dt.timedelta(days=1)))
        out.append((dt.datetime(y, m, d2), dt.datetime(nxt_y, nxt_m, 14)))
        y, m = nxt_y, nxt_m
    return out

pay_dates = semi_monthly(dt.date(2026, 8, 1), (2028, 8))
PP0 = PB + 2
for i, (pdt, pend) in enumerate(pay_dates):
    r = PP0 + i
    rng = f'{TXD},">="&$A{r},{TXD},"<="&$L{r}'
    cell(mo, r, 1, pdt, DATEF, align="center")
    cell(mo, r, 2, f'=TEXT($A{r},"mmm d")&" - "&TEXT($L{r},"mmm d")', align="center")
    cell(mo, r, 3, f'=SUMIFS({TXIN},{rng}{NT})', MONEY)
    cell(mo, r, 4, f'=SUMIFS({TXOUT},{rng}{NT})', MONEY)
    cell(mo, r, 5, f'=SUMIFS({TXOUT},{rng},{TXC},"Student Debt")', MONEY)
    cell(mo, r, 6, f'=SUMIFS({TXOUT},{rng},{TXC},"Savings/Investments")', MONEY)
    cell(mo, r, 7, f"=$D{r}-$E{r}-$F{r}", MONEY)
    cell(mo, r, 8, f"=$C{r}-$D{r}", MONEY, font=F_BOLD)
    cell(mo, r, 9, f'=IF($C{r}>0,"yes","-")', align="center")
    mo.cell(row=r, column=12, value=pend).number_format = DATEF
mo.column_dimensions["L"].hidden = True
PP1 = PP0 + len(pay_dates) - 1
zebra(mo, f"A{PP0}:I{PP1}", PP0)
mo.conditional_formatting.add(f"H{PP0}:H{PP1}", FormulaRule(
    formula=[f'AND($C{PP0}>0,$H{PP0}<0)'],
    fill=PatternFill("solid", bgColor=BADF), font=Font(bold=True, color=RED)))

# =====================================================================
# 5. SAVINGS
# =====================================================================
sheet_title(sv, "Savings",
            "Where your savings are going. Log savings on Transactions with Category = "
            "Savings/Investments and pick the bucket. The only thing you set here is the "
            "target split.", 7)
for i, w in enumerate([22, 15, 13, 15, 16, 13, 15, 4, 4, 4, 4, 12]):
    sv.column_dimensions[get_column_letter(i + 1)].width = w

SAVED_ALL = f'SUMIFS({TXOUT},{TXC},"Savings/Investments")'
MONTHS_ELAPSED = 'MAX(1,(YEAR(TODAY())-YEAR($A$26))*12+MONTH(TODAY())-MONTH($A$26)+1)'

banner(sv, 4, 1, 7, "SAVINGS SNAPSHOT")
kpi(sv, 5, 1, "saved to date", "=" + SAVED_ALL, MONEY0)
kpi(sv, 5, 2, "this month", (
    f'=SUMIFS({TXOUT},{TXD},">="&EOMONTH(TODAY(),-1)+1,{TXD},"<="&EOMONTH(TODAY(),0),'
    f'{TXC},"Savings/Investments")'), MONEY0)
kpi(sv, 5, 3, "monthly average", "=" + SAVED_ALL + "/" + MONTHS_ELAPSED, MONEY0, small=True)
kpi(sv, 5, 4, "planned / month", f"={PLAN_SAVE}", MONEY0, small=True)
kpi(sv, 5, 5, "vs plan this month", f"=$B$6-{PLAN_SAVE}", MONEY0, small=True)
kpi(sv, 5, 6, "buckets in use", "=$D$21", INT, small=True)
kpi(sv, 5, 7, "biggest bucket", "=$A$21", TEXTF, small=True)
sv.cell(row=6, column=1).font = Font(bold=True, size=16, color=GREEN)
sv.conditional_formatting.add("E6", FormulaRule(
    formula=['$E$6<0'], font=Font(bold=True, size=12, color=RED)))
sv.conditional_formatting.add("E6", FormulaRule(
    formula=['$E$6>=0'], font=Font(bold=True, size=12, color=GREEN)))

banner(sv, 8, 1, 7, "WHERE IT GOES")
month_picker(sv, 9, 2, 12)
SSEL = "$B$9"
cell(sv, 9, 4, '="Contributions during "&TEXT($B$9,"mmmm yyyy")',
     font=Font(italic=True, size=10, color=GREY), border=NOBOX)
table_head(sv, 10, 1, ["Bucket", "All time", "Target split", "Planned / month",
                       "Selected month", "Actual split", "vs target / mo"],
           None, height=30)
SB0 = 11
SB1 = SB0 + len(SAV_BUCKETS) - 1
defaults = {"Masters Fund": 0.40, "Brokerage / Stocks": 0.35,
            "Emergency Fund": 0.15, "Roth IRA": 0.10}
for i, b in enumerate(SAV_BUCKETS):
    r = SB0 + i
    cell(sv, r, 1, b, font=F_BOLD)
    cell(sv, r, 2, f'=SUMIFS({TXOUT},{TXC},"Savings/Investments",{TXG},$A{r})', MONEY)
    cell(sv, r, 3, defaults.get(b, 0.0), PCT, fill=FILL_INPUT, align="center")
    cell(sv, r, 4, f"={PLAN_SAVE}*$C{r}", MONEY, fill=FILL_BAND)
    cell(sv, r, 5, "=" + m_out(SSEL, f',{TXC},"Savings/Investments",{TXG},$A{r}'), MONEY)
    cell(sv, r, 6, f'=IFERROR($B{r}/SUM($B${SB0}:$B${SB1}),0)', PCT, align="center")
    cell(sv, r, 7, f"=$E{r}-$D{r}", MONEY)
ST = SB1 + 1
cell(sv, ST, 1, "TOTAL", font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND)
for col in (2, 3, 4, 5, 6, 7):
    L = get_column_letter(col)
    cell(sv, ST, col, f"=SUM({L}{SB0}:{L}{SB1})",
         PCT if col in (3, 6) else MONEY,
         font=Font(bold=True, size=10, color=NAVY), fill=FILL_BAND, align="center")
zebra(sv, f"A{SB0}:G{SB1}", SB0)
sv.conditional_formatting.add(f"B{SB0}:B{SB1}", DataBarRule(
    start_type="num", start_value=0, end_type="max", color="7FBF9B", showValue=True))
sv.conditional_formatting.add(f"E{SB0}:E{SB1}", DataBarRule(
    start_type="num", start_value=0, end_type="max", color="9FC6E0", showValue=True))
sv.conditional_formatting.add(f"C{ST}", FormulaRule(
    formula=[f'ABS($C${ST}-1)>0.001'],
    fill=PatternFill("solid", bgColor=BADF), font=Font(bold=True, color=RED)))
cell(sv, ST + 1, 1, "Target split is yours to edit - the yellow cells. "
                    "It turns red if the total is not 100%.", font=F_NOTE, border=NOBOX)

cell(sv, 20, 1, "BIGGEST BUCKET", font=F_LBL, fill=FILL_CARD, align="center")
cell(sv, 20, 2, "AMOUNT", font=F_LBL, fill=FILL_CARD, align="center")
cell(sv, 20, 3, "SHARE", font=F_LBL, fill=FILL_CARD, align="center")
cell(sv, 20, 4, "BUCKETS IN USE", font=F_LBL, fill=FILL_CARD, align="center")
cell(sv, 21, 1, (f'=IFERROR(INDEX($A${SB0}:$A${SB1},'
                 f'MATCH(MAX($B${SB0}:$B${SB1}),$B${SB0}:$B${SB1},0)),"-")'),
     TEXTF, font=F_KPI_S, fill=FILL_CARD, align="center")
cell(sv, 21, 2, f"=MAX($B${SB0}:$B${SB1})", MONEY0, font=F_KPI_S, fill=FILL_CARD, align="center")
cell(sv, 21, 3, f'=IFERROR($B$21/$B${ST},0)', PCT, font=F_KPI_S, fill=FILL_CARD, align="center")
cell(sv, 21, 4, f'=COUNTIF($B${SB0}:$B${SB1},">0")', INT, font=F_KPI_S,
     fill=FILL_CARD, align="center")

banner(sv, 24, 1, 7, "SAVINGS BY MONTH")
table_head(sv, 25, 1, ["Month", "Contributed", "Cumulative", "Plan", "vs plan"],
           None, height=22)
SM0, SM1 = 26, 26 + N_MONTHS - 1
for i in range(N_MONTHS):
    r = SM0 + i
    cell(sv, r, 1, START if i == 0 else f"=EDATE($A{r-1},1)", MONF, font=F_BOLD, align="center")
    cell(sv, r, 2, "=" + m_out(f"$A{r}", ss), MONEY)
    cell(sv, r, 3, f'=$B{r}' if i == 0 else f'=$C{r-1}+$B{r}', MONEY, font=F_BOLD)
    cell(sv, r, 4, f"={PLAN_SAVE}", MONEY, fill=FILL_BAND)
    cell(sv, r, 5, f'=IF($B{r}=0,"",$B{r}-$D{r})', MONEY)
zebra(sv, f"A{SM0}:E{SM1}", SM0)
sv.conditional_formatting.add(f"E{SM0}:E{SM1}", FormulaRule(
    formula=[f'AND(ISNUMBER($E{SM0}),$E{SM0}<0)'],
    fill=PatternFill("solid", bgColor=BADF), font=Font(bold=True, color=RED)))
sv.conditional_formatting.add(f"E{SM0}:E{SM1}", FormulaRule(
    formula=[f'AND(ISNUMBER($E{SM0}),$E{SM0}>=0)'],
    fill=PatternFill("solid", bgColor=GOODF), font=Font(bold=True, color=GREEN)))

ch3 = LineChart()
ch3.title = "Savings piling up"
ch3.style = 2
ch3.height, ch3.width = 8, 15
ch3.y_axis.numFmt = MONEY0
ch3.add_data(Reference(sv, min_col=3, max_col=3, min_row=25, max_row=SM1),
             titles_from_data=True)
ch3.set_categories(Reference(sv, min_col=1, min_row=SM0, max_row=SM1))
sv.add_chart(ch3, "I10")

# =====================================================================
wb.calculation.fullCalcOnLoad = True
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.active = 0
wb.save(OUT)
print("saved:", OUT)
print("sheets:", wb.sheetnames)
